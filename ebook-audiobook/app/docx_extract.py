"""Extraction du texte d'un document Word (.docx) avec python-docx.

Le DOCX est beaucoup plus simple que le PDF : il contient de vrais
paragraphes et, souvent, des styles de titre (« Titre 1 », « Heading 1 »).
En-têtes, pieds de page et notes de bas de page sont stockés à part dans le
fichier : ils ne sont donc jamais lus. Il reste à retirer les appels de note
saisis à la main en exposant et les sauts de ligne manuels.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.text.hyperlink import Hyperlink
from docx.text.run import Run

from .cleaning import clean_text, is_scene_break
from .models import BREAK, HEADING, PARAGRAPH, Block, ExtractedDocument

# « Heading 1 », « Titre 2 », « Überschrift 1 »... : on récupère le chiffre.
_HEADING_STYLE = re.compile(r"(heading|titre|title|berschrift|título|titolo)\s*(\d)", re.I)
_NOTE_CALL = re.compile(r"^\s*[\d*†‡,]+\s*$")


def _heading_level(paragraph) -> int:
    """Niveau de titre du paragraphe (0 si ce n'est pas un titre)."""
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    m = _HEADING_STYLE.search(style)
    if m:
        return max(1, int(m.group(2)))
    if style.lower() in ("title", "titre"):
        return 1
    # Niveau hiérarchique défini directement sur le paragraphe.
    lvl = paragraph._p.xpath("./w:pPr/w:outlineLvl/@w:val")
    if lvl and lvl[0].isdigit() and int(lvl[0]) < 9:
        return int(lvl[0]) + 1
    return 0


def _paragraph_text(paragraph) -> str:
    """Texte du paragraphe sans les appels de note en exposant."""
    parts: list[str] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Run):
            text = item.text
            if item.font.superscript and _NOTE_CALL.match(text or ""):
                continue
            parts.append(text)
        elif isinstance(item, Hyperlink):
            parts.append(item.text)
    # Sauts de ligne manuels (Maj+Entrée) et tabulations -> espaces.
    return clean_text("".join(parts).replace("\n", " ").replace("\t", " "))


def _first_image(document) -> tuple[bytes | None, str]:
    """Première image JPEG/PNG du document, utilisée comme pochette."""
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            try:
                blob = rel.target_part.blob
            except Exception:
                continue
            if blob[:3] == b"\xff\xd8\xff":
                return blob, "jpg"
            if blob[:8] == b"\x89PNG\r\n\x1a\n":
                return blob, "png"
    return None, "jpg"


def extract_docx(path: str | Path, progress=None) -> ExtractedDocument:
    """Extrait un fichier .docx en blocs de texte structurés."""
    document = docx.Document(str(path))
    blocks: list[Block] = []
    paragraphs = document.paragraphs
    total = len(paragraphs)
    doc_title = ""

    for i, p in enumerate(paragraphs):
        text = _paragraph_text(p)
        if progress and i % 200 == 0:
            progress(i, total)
        if not text:
            continue
        style = (p.style.name if p.style is not None else "").lower()
        if style in ("title", "titre") and not doc_title:
            doc_title = text
        level = _heading_level(p)
        if level:
            blocks.append(Block(HEADING, text, level))
        elif is_scene_break(text):
            blocks.append(Block(BREAK, ""))
        else:
            blocks.append(Block(PARAGRAPH, text))

    props = document.core_properties
    cover, ext = _first_image(document)
    return ExtractedDocument(
        title=clean_text(props.title or "") or doc_title or Path(path).stem,
        author=clean_text(props.author or ""),
        blocks=blocks,
        cover=cover,
        cover_ext=ext,
    )
